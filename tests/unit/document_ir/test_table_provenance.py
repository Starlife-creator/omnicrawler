from __future__ import annotations

import hashlib
import json

from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from omnicrawler.document_ir import parse_document
from omnicrawler.services.archive_analysis import execute


def test_actual_pdf_table_cells_and_regions_reach_archive_evidence(tmp_path):
    path = tmp_path / "table.pdf"
    table = Table([["Item", "Price"], ["Paper", "12.50"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
    SimpleDocTemplate(str(path)).build([table])
    parsed = parse_document(path)
    assert parsed.tables == [[["Item", "Price"], ["Paper", "12.50"]]]
    assert parsed.table_locators[0]["page"] == 1
    assert len(parsed.table_locators[0]["bbox"]) == 4
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"format": 1, "sources": [{"id": "table", "path": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}]}))
    result = execute(manifest, tmp_path / "report")
    report = json.loads(__import__("pathlib").Path(result["report"]).read_text(encoding="utf-8"))
    row = next(item for item in report["evidence"] if item["locator"].get("row") == 2)
    assert row["quote"] == "Paper | 12.50"
    assert row["locator"]["page"] == 1 and row["locator"]["bbox"] == parsed.table_locators[0]["bbox"]


def test_empty_docx_can_parse_without_indexing_empty_paragraphs(tmp_path):
    from docx import Document
    path = tmp_path / "empty.docx"
    Document().save(path)
    assert parse_document(path).paragraphs == []
