"""Export must preserve the order and source of interleaved document blocks."""
from omnicrawler.document_ir import parse_document


def test_docx_interleaved_blocks_and_table_header_escaping(tmp_path):
    import docx
    source = docx.Document()
    source.add_heading("Report", level=1)
    source.add_paragraph("before")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A|B"
    table.cell(0, 1).text = "C\nD"
    table.cell(1, 0).text = "data"
    table.cell(1, 1).text = "value"
    source.add_heading("After table", level=2)
    source.add_paragraph("after")
    path = tmp_path / "ordered.docx"
    source.save(path)
    result = parse_document(path)
    for view in (result.to_text(), result.to_markdown()):
        assert view.index("before") < view.index("data") < view.index("after")
    markdown = result.to_markdown()
    assert "A\\|B" in markdown
    assert "C D" in markdown
    assert "## After table" in markdown
    assert result.table_locators[0]["body_index"] == 2
    assert result.paragraph_locators[-1]["body_index"] == 4


def test_html_repeated_facts_are_retained_in_source_order(tmp_path):
    path = tmp_path / "table.html"
    path.write_text("<main><h1>Title</h1><p>repeated</p><table><tr><th>key</th></tr><tr><td>value</td></tr></table><h2>Section</h2><p>repeated</p></main>", encoding="utf-8")
    result = parse_document(path)
    assert result.paragraphs.count("repeated") == 2
    assert result.tables == [[["key"], ["value"]]]
    view = result.to_markdown()
    assert view.index("repeated") < view.index("value") < view.rindex("repeated")
    assert "## Section" in view
