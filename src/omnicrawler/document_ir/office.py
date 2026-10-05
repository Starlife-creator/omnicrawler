"""document_ir 富文档解析（S2）：.docx / .pptx / .odt / .epub。

依赖策略：
- .docx → python-docx（懒加载，缺失时抛 ModuleNotFoundError 提示 omnicrawler[document]）
- .pptx → python-pptx（懒加载，同上）
- .odt  → zipfile + xml.etree.ElementTree（标准库，零依赖）
- .epub → zipfile + 项目内 html_tools（标准库，零依赖）

注册方式与 parsers.py 一致：register_document_parser 装饰器。
"""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

from ..extraction.html_tools import parse_html
from .base import DocumentIR
from .parsers import _append_html_blocks, _require_file, _require_parser, register_document_parser

#: 正文候选容器（与 parsers.py 的 HTML 解析一致）
_CONTENT_SELECTOR = "p,li,blockquote,h1,h2,h3,h4,h5,h6"

#: ODT 文本命名空间（段落标签）
_ODT_TEXT_NS = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def _dedupe(items: list[str]) -> list[str]:
    """顺序去重（HTML 各节点文本可能重叠）。"""
    out: list[str] = []
    for item in items:
        if item and item not in out:
            out.append(item)
    return out


# ── .docx ────────────────────────────────────────────────
@register_document_parser(".docx")
def _parse_docx(path: Path, options: dict[str, Any]) -> DocumentIR:
    _require_file(path)
    try:
        import docx  # python-docx
    except ImportError:
        _require_parser(".docx", "python-docx")

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    core = document.core_properties
    result = DocumentIR(source=path, kind=".docx", title=core.title or path.stem,
                        metadata={"author": core.author or ""})
    title_consumed = False
    for index, element in enumerate(document.element.body):
        locator = {"part": "word/document.xml", "body_index": index}
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            if not title_consumed:
                result.title = text
                title_consumed = True
                result.metadata["title_locator"] = locator
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            level = int(style[-1]) if style.startswith("Heading ") and style[-1:].isdigit() else 0
            result.add_paragraph(text, locator=locator, heading_level=level)
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            result.add_table([[cell.text.strip() for cell in row.cells] for row in table.rows], locator=locator)
    if core.created is not None:
        result.metadata["created"] = core.created.isoformat()
    if core.modified is not None:
        result.metadata["modified"] = core.modified.isoformat()
    result.metadata["paragraph_count"] = len(result.paragraphs)
    return result


# ── .pptx ────────────────────────────────────────────────
@register_document_parser(".pptx")
def _parse_pptx(path: Path, options: dict[str, Any]) -> DocumentIR:
    _require_file(path)
    try:
        from pptx import Presentation
    except ImportError:
        _require_parser(".pptx", "python-pptx")

    prs = Presentation(str(path))
    result = DocumentIR(source=path, kind=".pptx", title=path.stem)
    for slide_no, slide in enumerate(prs.slides, 1):
        for shape_no, shape in enumerate(slide.shapes, 1):
            locator = {"slide": slide_no, "shape": shape_no, "shape_id": shape.shape_id,
                       "order_source": "shape_tree"}
            if shape.has_text_frame:
                for paragraph_no, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        result.add_paragraph(text, locator={**locator, "paragraph": paragraph_no})
            elif shape.has_table:
                result.add_table([[cell.text.strip() for cell in row.cells] for row in shape.table.rows], locator=locator)
    result.promote_first_paragraph_to_title()
    result.metadata.update(slide_count=len(prs.slides), paragraph_count=len(result.paragraphs))
    return result


# ── .odt ─────────────────────────────────────────────────
@register_document_parser(".odt")
def _parse_odt(path: Path, options: dict[str, Any]) -> DocumentIR:
    _require_file(path)
    try:
        with zipfile.ZipFile(path) as zf:
            content = _read_zip_entry(zf, "content.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError(f"document_ir: 无效的 ODT 文件: {path} ({exc})") from exc

    tree = ET.parse(io.BytesIO(content))
    result = DocumentIR(source=path, kind=".odt", title=path.stem)
    table_ns = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
    def visit(element: Any, location: str) -> None:
        if element.tag == table_ns + "table":
            cells = [[" ".join("".join(cell.itertext()).split()) for cell in row
                      if cell.tag in {table_ns + "table-cell", table_ns + "covered-table-cell"}]
                     for row in element.iter(table_ns + "table-row")]
            result.add_table(cells, locator={"part": "content.xml", "element_path": location})
        elif element.tag in {_ODT_TEXT_NS + "p", _ODT_TEXT_NS + "h"}:
            text = " ".join("".join(element.itertext()).split())
            if text:
                level = element.get(_ODT_TEXT_NS + "outline-level", "0")
                result.add_paragraph(text, locator={"part": "content.xml", "element_path": location},
                                     heading_level=int(level) if level.isdigit() else 0)
        else:
            for index, child in enumerate(element):
                visit(child, f"{location}/{index}")
    visit(tree.getroot(), "0")
    result.promote_first_paragraph_to_title()
    result.metadata["paragraph_count"] = len(result.paragraphs)
    return result


# ── .epub ────────────────────────────────────────────────
# B07-002：zip 单条目解压上限（64 MiB）——防 zip 炸弹/超大 content.xml 内存耗尽
_ZIP_ENTRY_MAX_BYTES = 64 * 1024 * 1024


def _read_zip_entry(
    zf: zipfile.ZipFile, name: str, *, max_bytes: int = _ZIP_ENTRY_MAX_BYTES,
) -> bytes:
    """解压 zip 条目前检查 file_size 上限（defense-in-depth）。"""
    info = zf.getinfo(name)
    if info.file_size > max_bytes:
        raise ValueError(
            f"document_ir: zip 条目过大（{info.file_size} 字节 > 上限 {max_bytes}）: {name}"
        )
    return zf.read(name)


def _epub_spine_order(zf: zipfile.ZipFile) -> list[str]:
    """按 EPUB 阅读顺序返回内容文件相对路径列表；无 OPF 时返回空。"""
    try:
        container = ET.parse(io.BytesIO(_read_zip_entry(zf, "META-INF/container.xml")))
    except (KeyError, ET.ParseError):
        return []
    ns = {"c": "urn:oasis:names:tc:opendocument:xmlns:container"}
    rootfile = container.find(".//c:rootfile", ns)
    if rootfile is None:
        return []
    opf_path = rootfile.get("full-path")
    if not opf_path:
        return []
    try:
        opf = ET.parse(io.BytesIO(_read_zip_entry(zf, opf_path)))
    except (KeyError, ET.ParseError):
        return []

    base_dir = opf_path.rsplit("/", 1)[0] if "/" in opf_path else ""
    manifest: dict[str, str] = {}
    for item in opf.findall(".//{http://www.idpf.org/2007/opf}item"):
        item_id = item.get("id")
        href = item.get("href")
        if item_id and href:
            manifest[item_id] = f"{base_dir}/{href}" if base_dir else href
    order: list[str] = []
    for ref in opf.findall(".//{http://www.idpf.org/2007/opf}itemref"):
        idref = ref.get("idref")
        if idref and idref in manifest:
            order.append(manifest[idref])
    return order


@register_document_parser(".epub")
def _parse_epub(path: Path, options: dict[str, Any]) -> DocumentIR:
    _require_file(path)
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ValueError(f"document_ir: 无效的 EPUB 文件: {path} ({exc})") from exc

    with zf:
        content_files = [
            name
            for name in zf.namelist()
            if name.lower().endswith((".xhtml", ".html", ".htm"))
        ]
        spine = _epub_spine_order(zf)
        ordered = [name for name in spine if name in content_files]
        if not ordered:
            ordered = sorted(content_files)

        result = DocumentIR(source=path, kind=".epub", title=path.stem)
        for name in ordered:
            raw = _read_zip_entry(zf, name)
            document = parse_html(raw.decode("utf-8", errors="replace"))
            _append_html_blocks(result, document, part=name, include_h1=True)
    result.promote_first_paragraph_to_title()
    result.metadata.update(content_files=len(ordered), paragraph_count=len(result.paragraphs))
    return result


__all__ = ["_epub_spine_order"]
