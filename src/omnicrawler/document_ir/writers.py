"""DocumentIR 出口：HTML / DOCX / PDF / XLSX / 文本 / Markdown。

设计边界：
- 只依赖 ``DocumentIR``，不依赖 convertx（避免包级循环）。
- 富文档写出按需懒加载：python-docx / reportlab / openpyxl 缺失时抛
  ``ModuleNotFoundError``，由上层给出安装提示。
- 文本类写出走 :func:`core.utils.atomic_write_text`，xlsx 内部临时文件 + 原子替换（B-P3-5）。
- 表格与标题层级是 IR 的一等公民：docx/pdf/xlsx 导出保留表格，而不是退化为纯文本。
"""

from __future__ import annotations

import html as _html
import logging
import os
import re
import tempfile
import xml.sax.saxutils as _xml_escape
from pathlib import Path
from typing import Any

from ..core.utils import atomic_write_text, excel_safe
from .base import DocumentIR

LOGGER = logging.getLogger(__name__)

_MAX_SHEET_TITLE = 31
_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _paragraph_text(ir: DocumentIR, index: int) -> str:
    if 0 <= index < len(ir.paragraphs):
        return ir.paragraphs[index]
    return ""


def _table_rows(table: list[list[str]]) -> tuple[list[list[str]], int]:
    width = max((len(row) for row in table), default=0)
    if not width:
        return [], 0
    return [list(row) + [""] * (width - len(row)) for row in table], width


def render_html(ir: DocumentIR) -> str:
    """把 IR 渲染为独立 HTML5 文档（含 UTF-8 meta 与表格）。"""
    title = _html.escape(ir.title or ir.source.stem)
    lines = [
        "<!doctype html>",
        '<html lang="zh-CN">',
        "<head>",
        '  <meta charset="utf-8">',
        f"  <title>{title}</title>",
        "  <style>",
        "    body { font-family: system-ui, -apple-system, 'Segoe UI', sans-serif;"
        " max-width: 860px; margin: 2rem auto; line-height: 1.65; padding: 0 1rem; }",
        "    table { border-collapse: collapse; width: 100%; margin: 1rem 0; }",
        "    th, td { border: 1px solid #bbb; padding: .35rem .5rem; text-align: left;"
        " vertical-align: top; }",
        "    th { background: #f2f2f2; }",
        "  </style>",
        "</head>",
        "<body>",
        "<article>",
    ]
    if ir.title:
        lines.append(f"  <h1>{_html.escape(ir.title)}</h1>")
    for block in ir.ordered_blocks():
        if block.kind == "paragraph":
            text = _paragraph_text(ir, block.index)
            if not text:
                continue
            if block.heading_level:
                level = min(6, max(1, int(block.heading_level)))
                lines.append(f"  <h{level}>{_html.escape(text)}</h{level}>")
            else:
                lines.append(f"  <p>{_html.escape(text)}</p>")
            continue
        rows, _width = _table_rows(ir.tables[block.index])
        if not rows:
            continue
        lines.append("  <table>")
        lines.append("    <thead><tr>")
        lines.extend(f"      <th>{_html.escape(cell)}</th>" for cell in rows[0])
        lines.append("    </tr></thead>")
        lines.append("    <tbody>")
        for row in rows[1:]:
            lines.append("      <tr>" + "".join(f"<td>{_html.escape(cell)}</td>" for cell in row) + "</tr>")
        lines.append("    </tbody>")
        lines.append("  </table>")
    lines.extend(["</article>", "</body>", "</html>", ""])
    return "\n".join(lines)


def write_text(ir: DocumentIR, path: Path, options: dict[str, Any]) -> dict[str, Any]:
    content = ir.to_text()
    atomic_write_text(path, content, encoding=str(options.get("encoding", "utf-8")))
    return {"rows": ir.paragraph_count() + ir.table_count(), "chars": len(content), "tables": ir.table_count()}


def write_markdown(ir: DocumentIR, path: Path, options: dict[str, Any]) -> dict[str, Any]:
    content = ir.to_markdown()
    atomic_write_text(path, content, encoding=str(options.get("encoding", "utf-8")))
    return {"rows": ir.paragraph_count() + ir.table_count(), "chars": len(content), "tables": ir.table_count()}


def write_html(ir: DocumentIR, path: Path, options: dict[str, Any]) -> dict[str, Any]:
    content = render_html(ir)
    atomic_write_text(path, content, encoding=str(options.get("encoding", "utf-8")))
    return {"rows": ir.paragraph_count() + ir.table_count(), "chars": len(content), "tables": ir.table_count()}


def write_docx(ir: DocumentIR, path: Path, options: dict[str, Any]) -> dict[str, Any]:
    try:
        import docx
    except ModuleNotFoundError as exc:  # pragma: no cover - 依赖缺失路径
        raise ModuleNotFoundError(
            "写出 .docx 需要 python-docx：pip install 'omnicrawler[document]'", name="docx"
        ) from exc

    document = docx.Document()
    if ir.title:
        document.add_heading(ir.title, level=0)
    for block in ir.ordered_blocks():
        if block.kind == "paragraph":
            text = _paragraph_text(ir, block.index)
            if not text:
                continue
            if block.heading_level:
                document.add_heading(text, level=min(9, max(1, int(block.heading_level))))
            else:
                document.add_paragraph(text)
            continue
        rows, width = _table_rows(ir.tables[block.index])
        if not rows:
            continue
        table = document.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index, cell_text in enumerate(row):
                table.cell(row_index, column_index).text = cell_text
    try:
        document.core_properties.title = ir.title
    except Exception:  # noqa: BLE001 - 旧版 python-docx 无 core_properties 时不影响正文
        # B-P3-3：元数据写入失败留痕，不再静默吞掉。
        LOGGER.debug("DOCX core_properties 标题写入失败，跳过元数据: ", exc_info=True)
    document.save(str(path))
    return {"rows": ir.paragraph_count() + ir.table_count(), "tables": ir.table_count(), "format": "docx"}


_PDF_FONT: str | None = None


def _pdf_font() -> str:
    global _PDF_FONT
    if _PDF_FONT is not None:
        return _PDF_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        _PDF_FONT = "STSong-Light"
    except Exception:  # noqa: BLE001 - 无 CJK 字体时退回内置字体，正文仍可写出
        _PDF_FONT = "Helvetica"
    return _PDF_FONT


def write_pdf(ir: DocumentIR, path: Path, options: dict[str, Any]) -> dict[str, Any]:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ModuleNotFoundError as exc:  # pragma: no cover - 依赖缺失路径
        raise ModuleNotFoundError(
            "写出 .pdf 需要 reportlab：pip install 'omnicrawler[pdf]'", name="reportlab"
        ) from exc

    font = _pdf_font()
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=16)
    styles = {
        1: ParagraphStyle("h1", parent=body, fontSize=18, leading=24, spaceAfter=8),
        2: ParagraphStyle("h2", parent=body, fontSize=15, leading=21, spaceAfter=6),
        3: ParagraphStyle("h3", parent=body, fontSize=12.5, leading=18, spaceAfter=4),
    }

    def paragraph(text: str, style: Any = body) -> Any:
        return Paragraph(_xml_escape.escape(text).replace("\n", "<br/>"), style)

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=ir.title or ir.source.stem,
    )
    flowables: list[Any] = []
    if ir.title:
        flowables.append(paragraph(ir.title, styles[1]))
        flowables.append(Spacer(1, 4))
    available_width = A4[0] - 36 * mm
    for block in ir.ordered_blocks():
        if block.kind == "paragraph":
            text = _paragraph_text(ir, block.index)
            if not text:
                continue
            level = min(3, max(1, int(block.heading_level))) if block.heading_level else 0
            flowables.append(paragraph(text, styles.get(level, body)))
            flowables.append(Spacer(1, 3))
            continue
        rows, width = _table_rows(ir.tables[block.index])
        if not rows:
            continue
        cell_style = ParagraphStyle("cell", parent=body, fontSize=9, leading=12)
        data = [[paragraph(cell, cell_style) for cell in row] for row in rows]
        table = Table(data, colWidths=[available_width / width] * width, repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f2f2f2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        flowables.append(table)
        flowables.append(Spacer(1, 8))
    if not flowables:
        flowables.append(paragraph(ir.source.stem, body))
    document.build(flowables)
    return {"rows": ir.paragraph_count() + ir.table_count(), "tables": ir.table_count(), "format": "pdf"}


def _safe_sheet_title(value: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", value).strip() or "Sheet"
    base = base[:_MAX_SHEET_TITLE]
    candidate = base
    index = 2
    while candidate.lower() in used:
        suffix = f" ({index})"
        candidate = base[: _MAX_SHEET_TITLE - len(suffix)] + suffix
        index += 1
    used.add(candidate.lower())
    return candidate


def write_xlsx(ir: DocumentIR, path: Path, options: dict[str, Any]) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font
    except ModuleNotFoundError as exc:  # pragma: no cover - 依赖缺失路径
        raise ModuleNotFoundError(
            "写出 .xlsx 需要 openpyxl：pip install 'omnicrawler[pdf]' 或 'omnicrawler[gui]'",
            name="openpyxl",
        ) from exc

    # B-P1-1：标题/段落/表头/单元格统一经 excel_safe 中和公式注入；
    # B-P3-5：write_only 流式工作簿降低大文档内存峰值；原子替换收进函数内部。
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    temporary = Path(tmp_name)
    try:
        workbook = Workbook(write_only=True)
        default = workbook.create_sheet("Document")
        title_cell = WriteOnlyCell(default, value=excel_safe(ir.title or ir.source.stem))
        title_cell.font = Font(bold=True, size=13)
        default.append([title_cell])
        default.append([])
        for block in ir.ordered_blocks():
            if block.kind != "paragraph":
                continue
            text = _paragraph_text(ir, block.index)
            if not text:
                continue
            default.append([excel_safe(text)])
        used_titles: set[str] = {"document"}
        table_count = 0
        for table_index, table in enumerate(ir.tables, start=1):
            rows, width = _table_rows(table)
            if not rows:
                continue
            sheet = workbook.create_sheet(_safe_sheet_title(f"Table {table_index}", used_titles))
            for row in rows:
                sheet.append([excel_safe(cell_text) for cell_text in row])
            table_count += 1
        workbook.save(str(temporary))
        temporary.replace(path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return {"rows": ir.paragraph_count() + ir.table_count(), "tables": table_count, "format": "xlsx"}


#: 后缀 → 写出函数（不包含 xlsx：表格式输出由 convertx 的表 writer 负责）。
WRITERS: dict[str, Any] = {
    ".txt": write_text,
    ".md": write_markdown,
    ".html": write_html,
    ".htm": write_html,
    ".docx": write_docx,
    ".pdf": write_pdf,
}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        import json

        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def ir_from_records(
    rows: list[dict[str, Any]],
    *,
    title: str = "",
    source: Path | None = None,
) -> DocumentIR:
    """把表格记录提升为 DocumentIR：单文档记录按段落展开，其余按表格展开。"""
    source_path = source or Path("records")
    document = DocumentIR(source=source_path, kind=".records", title=title or source_path.stem)
    if not rows:
        return document
    single = rows[0]
    if len(rows) == 1 and "text" in single and ("title" in single or "record_type" in single):
        document.title = str(single.get("title") or document.title)
        text = str(single.get("text") or "")
        for paragraph_text in re.split(r"\n\s*\n", text):
            if paragraph_text.strip():
                document.add_paragraph(paragraph_text.strip())
        return document

    columns: dict[str, None] = {}
    for row in rows:
        for key in row:
            columns[str(key)] = None
    column_names = list(columns)
    table: list[list[str]] = [column_names]
    table.extend([[_cell_text(row.get(column, "")) for column in column_names] for row in rows])
    document.add_table(table)
    return document


def write_document_ir(ir: DocumentIR, path: Path, options: dict[str, Any]) -> dict[str, Any]:
    """按目标后缀写出 DocumentIR；不支持的后缀抛 ``ValueError``。"""
    suffix = path.suffix.lower()
    writer = WRITERS.get(suffix)
    if writer is None:
        raise ValueError(f"document_ir: 不支持写出格式 {suffix or '(无扩展名)'}")
    return dict(writer(ir, path, options))


__all__ = [
    "WRITERS",
    "ir_from_records",
    "render_html",
    "write_docx",
    "write_document_ir",
    "write_html",
    "write_markdown",
    "write_pdf",
    "write_text",
    "write_xlsx",
]
