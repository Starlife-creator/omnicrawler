"""Measure ConvertX text-format paths in isolated Python processes.

The report is informational: machine-dependent timing is deliberately not a
pass/fail gate.  Output correctness is checked by every worker before a sample
is accepted.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import platform
import subprocess
import sys
import tempfile
import time
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
CASE_FORMATS = {
    "csv-jsonl": ("csv", "jsonl"),
    "csv-jsonl-auto": ("csv", "jsonl"),
    "jsonl-csv": ("jsonl", "csv"),
    "jsonl-jsonl": ("jsonl", "jsonl"),
    "jsonl-xlsx": ("jsonl", "xlsx"),
    "xlsx-jsonl": ("xlsx", "jsonl"),
    "parquet-jsonl": ("parquet", "jsonl"),
    "duckdb-jsonl": ("duckdb", "jsonl"),
}
CASES = tuple(CASE_FORMATS)


def _record(index: int, rows: int) -> dict[str, str]:
    return {
        "record_id": f"r-{index:09d}",
        "source_url": f"https://example.invalid/items/{index}",
        "record_type": "benchmark",
        "created_at": "2026-09-06T00:00:00Z",
        "title": f"Item {index}",
        "note": f"line one\nline two {index}" if index % 997 == 0 else f"note {index}",
        "tail_only": "present" if index == rows - 1 else "",
    }


def generate_fixture(path: Path, *, format_name: str, rows: int) -> None:
    """Generate deterministic input outside the measured worker process."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if format_name == "csv":
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(_record(0, rows)))
            writer.writeheader()
            for index in range(rows):
                writer.writerow(_record(index, rows))
        return
    if format_name == "jsonl":
        with path.open("w", encoding="utf-8") as handle:
            for index in range(rows):
                handle.write(json.dumps(_record(index, rows), ensure_ascii=False) + "\n")
        return
    if format_name == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet()
        sheet.append(list(_record(0, rows)))
        for index in range(rows):
            sheet.append(list(_record(index, rows).values()))
        workbook.save(path)
        workbook.close()
        return
    if format_name == "parquet":
        import pyarrow as pa
        import pyarrow.parquet as pq

        parquet_writer: Any = None
        try:
            for start in range(0, rows, 10_000):
                batch = [_record(index, rows) for index in range(start, min(rows, start + 10_000))]
                table = pa.Table.from_pylist(batch)
                if parquet_writer is None:
                    parquet_writer = pq.ParquetWriter(path, table.schema, compression="zstd")
                parquet_writer.write_table(table)
        finally:
            if parquet_writer is not None:
                parquet_writer.close()
        return
    if format_name == "duckdb":
        import duckdb

        connection = duckdb.connect(str(path))
        try:
            connection.execute(
                "CREATE TABLE records (record_id VARCHAR, source_url VARCHAR, record_type VARCHAR, "
                "created_at VARCHAR, title VARCHAR, note VARCHAR, tail_only VARCHAR)"
            )
            for start in range(0, rows, 10_000):
                batch = [_record(index, rows) for index in range(start, min(rows, start + 10_000))]
                connection.executemany("INSERT INTO records VALUES (?, ?, ?, ?, ?, ?, ?)", [list(row.values()) for row in batch])
        finally:
            connection.close()
        return
    raise ValueError(f"unknown fixture format: {format_name}")


def _peak_rss_bytes() -> int | None:
    """Return the process peak resident set with no benchmark-only dependency."""
    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes

        class ProcessMemoryCounters(ctypes.Structure):
            _fields_ = [
                ("cb", wintypes.DWORD),
                ("PageFaultCount", wintypes.DWORD),
                ("PeakWorkingSetSize", ctypes.c_size_t),
                ("WorkingSetSize", ctypes.c_size_t),
                ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
                ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                ("PagefileUsage", ctypes.c_size_t),
                ("PeakPagefileUsage", ctypes.c_size_t),
            ]

        counters = ProcessMemoryCounters()
        counters.cb = ctypes.sizeof(counters)
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        psapi = ctypes.WinDLL("psapi", use_last_error=True)
        kernel32.GetCurrentProcess.restype = wintypes.HANDLE
        psapi.GetProcessMemoryInfo.argtypes = [
            wintypes.HANDLE,
            ctypes.POINTER(ProcessMemoryCounters),
            wintypes.DWORD,
        ]
        psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
        if not psapi.GetProcessMemoryInfo(
            kernel32.GetCurrentProcess(), ctypes.byref(counters), counters.cb
        ):
            return None
        return int(counters.PeakWorkingSetSize)
    try:
        import resource

        peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        return int(peak if sys.platform == "darwin" else peak * 1024)
    except (ImportError, OSError):
        return None


def _validate_output(path: Path, *, format_name: str, rows: int) -> dict[str, Any]:
    count = 0
    first: dict[str, Any] | None = None
    last: dict[str, Any] | None = None
    if format_name == "jsonl":
        with path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                value = json.loads(line)
                first = first or value
                last = value
                count += 1
    elif format_name == "csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for value in csv.DictReader(handle):
                first = first or value
                last = value
                count += 1
    elif format_name == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = workbook.active.iter_rows(values_only=True)
            header = [str(value) for value in next(values)]
            for cells in values:
                value = dict(zip(header, cells, strict=False))
                first = first or value
                last = value
                count += 1
        finally:
            workbook.close()
    else:
        raise ValueError(f"unknown output format: {format_name}")
    if count != rows:
        raise RuntimeError(f"expected {rows} output rows, got {count}")
    if rows:
        assert first is not None and last is not None
        if first.get("record_id") != "r-000000000":
            raise RuntimeError("first output record changed")
        if last.get("record_id") != f"r-{rows - 1:09d}":
            raise RuntimeError("last output record changed")
        if last.get("tail_only") != "present":
            raise RuntimeError("late-field value was lost")
    return {"rows": count, "first_and_last_match": True, "late_field_preserved": True}


def run_worker(case: str, source: Path, target: Path, rows: int) -> dict[str, Any]:
    """Run one measured conversion; called only in a fresh subprocess."""
    from omnicrawler.convertx import convert

    source_format, target_format = CASE_FORMATS[case]
    started = time.perf_counter()
    options = {"reader_csv": {"encoding": "auto"}} if case == "csv-jsonl-auto" else None
    result = convert(source, target, on_error="abort", options=options)
    elapsed = time.perf_counter() - started
    peak_rss = _peak_rss_bytes()
    correctness = _validate_output(target, format_name=target_format, rows=rows)
    if result.rows != rows or result.extra.get("written_records") != rows:
        raise RuntimeError("ConvertX result counts disagree with the generated workload")
    return {
        "case": case,
        "rows": rows,
        "duration_seconds": round(elapsed, 6),
        "rows_per_second": round(rows / elapsed, 2) if elapsed else None,
        "peak_rss_bytes": peak_rss,
        "input_bytes": source.stat().st_size,
        "output_bytes": target.stat().st_size,
        "correctness": correctness,
    }


def _git_revision() -> str | None:
    completed = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True, capture_output=True, check=False
    )
    return completed.stdout.strip() if completed.returncode == 0 else None


def _dependency_versions() -> dict[str, str]:
    found: dict[str, str] = {}
    for package in ("omnicrawler-platform", "PyYAML", "defusedxml"):
        try:
            found[package] = version(package)
        except PackageNotFoundError:
            found[package] = "not-installed"
    return found


def run_suite(
    *, sizes: list[int], repeats: int, work_dir: Path, cases: list[str] | None = None
) -> dict[str, Any]:
    selected_cases = list(cases or CASES)
    samples: list[dict[str, Any]] = []
    for rows in sizes:
        source_formats = {CASE_FORMATS[case][0] for case in selected_cases}
        inputs = {name: work_dir / f"input-{rows}.{name}" for name in source_formats}
        for name, path in inputs.items():
            generate_fixture(path, format_name=name, rows=rows)
        for case in selected_cases:
            source_format, target_format = CASE_FORMATS[case]
            for repeat in range(1, repeats + 1):
                target = work_dir / f"output-{case}-{rows}-{repeat}.{target_format}"
                command = [
                    sys.executable,
                    str(Path(__file__).resolve()),
                    "--worker",
                    "--case", case,
                    "--rows", str(rows),
                    "--source", str(inputs[source_format]),
                    "--target", str(target),
                ]
                completed = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, check=False)
                if completed.returncode != 0:
                    raise RuntimeError(
                        f"benchmark worker failed for {case}/{rows}: {completed.stderr.strip()}"
                    )
                sample = json.loads(completed.stdout)
                sample["repeat"] = repeat
                samples.append(sample)
    return {
        "schema_version": 1,
        "run_type": "fresh-process",
        "environment": {
            "python": platform.python_version(),
            "implementation": platform.python_implementation(),
            "platform": platform.platform(),
            "processor": platform.processor() or None,
            "logical_cpus": os.cpu_count(),
            "application_sha": _git_revision(),
            "dependencies": _dependency_versions(),
        },
        "parameters": {"sizes": sizes, "repeats": repeats, "cases": selected_cases},
        "samples": samples,
        "note": "Informational baseline; timing and RSS are not machine-independent pass/fail thresholds.",
    }


def _positive(value: str) -> int:
    number = int(value)
    if number <= 0:
        raise argparse.ArgumentTypeError("must be greater than zero")
    return number


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sizes", nargs="+", type=_positive, default=[10_000, 100_000, 1_000_000])
    parser.add_argument("--cases", nargs="+", choices=CASES, default=list(CASES))
    parser.add_argument("--repeats", type=_positive, default=1)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--work-dir", type=Path)
    parser.add_argument("--worker", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--case", choices=CASES, help=argparse.SUPPRESS)
    parser.add_argument("--rows", type=_positive, help=argparse.SUPPRESS)
    parser.add_argument("--source", type=Path, help=argparse.SUPPRESS)
    parser.add_argument("--target", type=Path, help=argparse.SUPPRESS)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.worker:
        if args.case is None or args.rows is None or args.source is None or args.target is None:
            raise SystemExit("worker mode requires --case, --rows, --source and --target")
        print(json.dumps(run_worker(args.case, args.source, args.target, args.rows)))
        return 0
    if args.work_dir is not None:
        args.work_dir.mkdir(parents=True, exist_ok=True)
        payload = run_suite(sizes=args.sizes, repeats=args.repeats, work_dir=args.work_dir, cases=args.cases)
    else:
        with tempfile.TemporaryDirectory(prefix="omnicrawler-convertx-benchmark-") as temp:
            payload = run_suite(sizes=args.sizes, repeats=args.repeats, work_dir=Path(temp), cases=args.cases)
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.output is not None:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
